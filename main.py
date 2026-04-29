# main.py - WattScope Final
from fastapi import FastAPI, Request, Depends, Form, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from sqlalchemy.orm import Session
from database import engine, Base, SessionLocal
from models import Client, ReleveQuotidien, Appareil
from analysis import get_dataframe_from_db, basic_stats, regression_temperature
import datetime
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

Base.metadata.create_all(bind=engine)

app = FastAPI(title="WattScope")

def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

@app.get("/", response_class=HTMLResponse)
async def accueil(request: Request, db: Session = Depends(get_db), success: str = ""):
    clients = db.query(Client).all()
    msg_success = f'<div class="alert alert-success">{success}</div>' if success else ""
    
    options = '<option value="">-- Choisissez --</option>'
    for c in clients:
        options += f'<option value="{c.id}">{c.nom_utilisateur} - {c.region}</option>'
    
    clients_html = ""
    for c in clients:
        nb = len(c.releves)
        clients_html += f"""<div style="background:rgba(255,215,0,0.2);border-left:4px solid #ffd700;padding:12px;margin-bottom:10px;border-radius:0 8px 8px 0"><div style="display:flex;justify-content:space-between;align-items:center"><div><strong>{c.nom_utilisateur}</strong> <span style="background:#6c757d;color:white;padding:2px 8px;border-radius:10px;margin-left:8px;font-size:.85em">{c.region}</span> <span style="color:#666;margin-left:8px">{c.type_logement}</span> <span class="badge bg-primary ms-2">{nb} relevés</span></div><a href="/dashboard/{c.id}" style="text-decoration:none;color:#333;border:1px solid #333;padding:4px 12px;border-radius:4px">📈 Dashboard</a></div></div>"""
    if not clients:
        clients_html = '<p style="color:#666">Aucun client.</p>'

    return HTMLResponse(content=f"""<!DOCTYPE html><html lang="fr"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width,initial-scale=1.0"><title>WattScope</title><link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet"><style>body{{background:linear-gradient(135deg,#0f2027,#203a43,#2c5364);min-height:100vh}}.card{{border-radius:15px;border:none}}.card-header{{border-radius:15px 15px 0 0!important;font-weight:bold}}.bg-gold{{background:linear-gradient(90deg,#ffd700,#ffaa00)}}.btn-primary{{background:#ffd700;border:none;color:#000;font-weight:bold}}.btn-success{{background:#28a745;border:none;color:white;font-weight:bold}}</style></head><body><div class="container py-5">{msg_success}<div class="text-center mb-5"><h1 class="display-4 text-white fw-bold">⚡ WattScope</h1><p class="lead text-light opacity-75">Suivi et Analyse de la Consommation Électrique</p></div><div class="row g-4"><div class="col-lg-7"><div class="card shadow mb-4"><div class="card-header bg-success text-white"><h2 class="h5 mb-0">➕ Nouveau Client</h2></div><div class="card-body"><form action="/ajouter_client" method="POST"><div class="row g-3"><div class="col-md-5"><input type="text" class="form-control" name="nom_utilisateur" placeholder="Ex: Jean-Yaoundé" required></div><div class="col-md-3"><input type="text" class="form-control" name="region" placeholder="Ville" required></div><div class="col-md-2"><select class="form-select" name="type_logement"><option>Studio</option><option>Appartement</option><option>Villa</option></select></div><div class="col-md-2"><button type="submit" class="btn btn-success w-100">➕ Ajouter</button></div></div></form></div></div><div class="card shadow"><div class="card-header bg-gold"><h2 class="h5 mb-0">📝 Nouveau Relevé</h2></div><div class="card-body"><form action="/ajouter_releve" method="POST"><div class="row g-3"><div class="col-md-6"><select class="form-select" name="foyer_id" required>{options}</select></div><div class="col-md-6"><input type="date" class="form-control" name="date_releve" required></div><div class="col-md-4"><input type="number" step="0.01" min="0" class="form-control" name="index_compteur" placeholder="kWh" required></div><div class="col-md-2"><select class="form-select" name="unite_coupure"><option value="minutes">Min</option><option value="heures">H</option></select></div><div class="col-md-2"><input type="number" min="0" step="0.1" class="form-control" name="duree_coupure" value="0"></div><div class="col-md-4"><input type="number" step="0.1" class="form-control" name="temperature" placeholder="°C (facultatif)"></div><div class="col-md-6"><input type="number" min="0" class="form-control" name="cout_estime" value="0" placeholder="FCFA"></div><div class="col-12"><button type="submit" class="btn btn-primary w-100 py-2">📊 Enregistrer</button></div></div></form></div></div><div class="card shadow mt-4"><div class="card-header bg-dark text-white"><h2 class="h5 mb-0">👥 Clients</h2></div><div class="card-body">{clients_html}</div></div></div><div class="col-lg-5"><div class="card shadow h-100"><div class="card-header bg-info text-white"><h2 class="h5 mb-0">💡 WattScope</h2></div><div class="card-body"><p>Analysez votre consommation électrique.</p><ul><li>📊 Régression linéaire</li><li>🔌 Appareils énergivores</li><li>📥 Export Excel</li></ul><hr><p class="small text-muted">TP INF 232 EC2</p></div></div></div></div></div><script src="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/js/bootstrap.bundle.min.js"></script></body></html>""")

@app.post("/ajouter_client")
async def ajouter_client(request: Request, nom_utilisateur: str = Form(...), region: str = Form(...), type_logement: str = Form("Appartement"), db: Session = Depends(get_db)):
    db.add(Client(nom_utilisateur=nom_utilisateur, region=region, type_logement=type_logement, nombre_habitants=1))
    db.commit()
    return RedirectResponse(url="/?success=✅ Client ajouté !", status_code=303)

@app.post("/ajouter_releve")
async def ajouter_releve(request: Request, foyer_id: int = Form(...), date_releve: str = Form(...), index_compteur: float = Form(...), duree_coupure: float = Form(0), unite_coupure: str = Form("minutes"), temperature: str = Form(""), cout_estime: float = Form(0.0), db: Session = Depends(get_db)):
    duree = int(duree_coupure * 60) if unite_coupure == "heures" else int(duree_coupure)
    temp = float(temperature) if temperature.strip() else None
    db.add(ReleveQuotidien(foyer_id=foyer_id, date_releve=datetime.date.fromisoformat(date_releve), index_compteur=index_compteur, duree_coupure_minutes=duree, temperature_exterieure=temp, cout_estime_fcfa=cout_estime))
    db.commit()
    return RedirectResponse(url=f"/dashboard/{foyer_id}?success=✅ Relevé enregistré", status_code=303)

@app.get("/dashboard/{foyer_id}", response_class=HTMLResponse)
async def dashboard(request: Request, foyer_id: int, db: Session = Depends(get_db), success: str = ""):
    client = db.query(Client).filter(Client.id == foyer_id).first()
    if not client: raise HTTPException(404, "Introuvable")
    msg = f'<div class="alert alert-success">{success}</div>' if success else ""
    df = get_dataframe_from_db(db)
    stats = basic_stats(df)
    reg = regression_temperature(df)
    releves = list(client.releves)
    conso_moy = round(sum(r.index_compteur for r in releves)/len(releves),2) if releves else 0

    # Appareils
    apps = db.query(Appareil).filter(Appareil.foyer_id == foyer_id).all()
    app_html = ""
    app_labels = "[]"
    app_data = "[]"
    top_app = ""
    if apps:
        total = sum(a.puissance_watts*a.heures_utilisation_jour*a.nombre_appareils/1000 for a in apps)
        tri = sorted(apps, key=lambda a: a.puissance_watts*a.heures_utilisation_jour*a.nombre_appareils, reverse=True)
        lbls, dts = [], []
        for a in tri:
            kwh = a.puissance_watts*a.heures_utilisation_jour*a.nombre_appareils/1000
            pct = round(kwh/total*100,1) if total>0 else 0
            badge = "🔴" if pct>30 else ("🟠" if pct>15 else "🟢")
            app_html += f"<tr><td>{badge} {a.nom_appareil}</td><td>{a.nombre_appareils}</td><td>{a.puissance_watts}W</td><td>{a.heures_utilisation_jour}h</td><td><strong>{round(kwh,2)}</strong></td><td><span class='badge {'bg-danger' if pct>30 else ('bg-warning' if pct>15 else 'bg-success')}'>{pct}%</span></td></tr>"
            lbls.append(a.nom_appareil); dts.append(round(kwh,2))
        top_app = f"<div class='alert alert-danger'><strong>⚠️ Énergivore:</strong> {tri[0].nom_appareil} ({round(tri[0].puissance_watts*tri[0].heures_utilisation_jour*tri[0].nombre_appareils/1000,2)} kWh/j)</div>"
        app_labels = "["+",".join([f"'{l}'" for l in lbls])+"]"
        app_data = "["+",".join([str(d) for d in dts])+"]"
    else:
        app_html = '<tr><td colspan="6" style="text-align:center;color:#666">Aucun</td></tr>'

    # Relevés
    rel_html = ""
    for r in releves[:20]:
        t = f"{r.temperature_exterieure}°C" if r.temperature_exterieure else "N/A"
        c = f"{r.cout_estime_fcfa} FCFA" if r.cout_estime_fcfa else "N/A"
        coup = f"{r.duree_coupure_minutes//60}h{r.duree_coupure_minutes%60:02d}" if r.duree_coupure_minutes>=60 else f"{r.duree_coupure_minutes} min"
        rel_html += f"<tr><td>{r.date_releve}</td><td><strong>{r.index_compteur}</strong></td><td>{coup}</td><td>{t}</td><td>{c}</td></tr>"
    if not releves: rel_html = '<tr><td colspan="5" style="text-align:center;color:#666">Aucun</td></tr>'

    reg_html = f"<p><strong>R²:</strong> {reg['r2_score']}</p><p>{reg['message']}</p>" if reg.get("r2_score") else f"<p>{reg.get('message','Données insuffisantes')}</p>"
    dates_js = "["+",".join([f"'{r.date_releve}'" for r in releves[:30]])+"]"
    consos_js = "["+",".join([str(r.index_compteur) for r in releves[:30]])+"]"

    return HTMLResponse(content=f"""<!DOCTYPE html><html lang="fr"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width,initial-scale=1.0"><title>Dashboard - {client.nom_utilisateur}</title><link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet"><script src="https://cdn.jsdelivr.net/npm/chart.js@4.4.0/dist/chart.umd.min.js"></script><style>body{{background:#f4f6f9}}.card{{border-radius:15px;border:none;margin-bottom:1.5rem}}.card-header{{border-radius:15px 15px 0 0!important;font-weight:bold}}.stat-box{{background:white;border-radius:12px;padding:1.25rem;text-align:center;box-shadow:0 2px 8px rgba(0,0,0,0.05)}}.stat-value{{font-size:2rem;font-weight:bold;color:#2c5364}}.stat-label{{font-size:.85rem;color:#666;text-transform:uppercase;letter-spacing:1px}}.btn-excel{{background:#217346;color:white;border:none;font-weight:bold}}</style></head><body><div class="container py-4">{msg}<div class="d-flex justify-content-between align-items-center mb-4"><div><h1 class="mb-0">⚡ Dashboard</h1><p class="text-muted mb-0">Client: <strong>{client.nom_utilisateur}</strong> | {client.region} | {client.type_logement}</p></div><div><a href="/export_excel/{foyer_id}" class="btn btn-excel me-2">📥 Excel</a><a href="/" class="btn btn-outline-secondary">← Retour</a></div></div><div class="row g-3 mb-4"><div class="col-md-3 col-6"><div class="stat-box"><div class="stat-value">{conso_moy}</div><div class="stat-label">kWh/jour</div></div></div><div class="col-md-3 col-6"><div class="stat-box"><div class="stat-value">{stats['total_releves']}</div><div class="stat-label">Total</div></div></div><div class="col-md-3 col-6"><div class="stat-box"><div class="stat-value">{stats['consommation_moyenne_kwh']}</div><div class="stat-label">kWh global</div></div></div><div class="col-md-3 col-6"><div class="stat-box"><div class="stat-value">{len(releves)}</div><div class="stat-label">Relevés</div></div></div><div class="row g-4 mb-4"><div class="col-lg-7"><div class="card shadow-sm"><div class="card-header bg-danger text-white">🔌 Appareils (kWh/j)</div><div class="card-body"><canvas id="appChart" height="200"></canvas></div></div></div><div class="col-lg-5"><div class="card shadow-sm"><div class="card-header bg-warning text-dark">⚠️ Énergivores</div><div class="card-body">{top_app}<div class="table-responsive"><table class="table table-sm"><thead><tr><th>Appareil</th><th>Qté</th><th>W</th><th>h/j</th><th>kWh</th><th>%</th></tr></thead><tbody>{app_html}</tbody></table></div><hr><p class="small text-muted">🔴>30% 🟠15-30% 🟢<15%</p><form action="/ajouter_appareil" method="POST"><input type="hidden" name="foyer_id" value="{foyer_id}"><div class="row g-2"><div class="col-5"><input type="text" class="form-control form-control-sm" name="nom_appareil" placeholder="Appareil" required></div><div class="col-3"><input type="number" class="form-control form-control-sm" name="puissance_watts" placeholder="Watts" required></div><div class="col-2"><input type="number" step="0.1" class="form-control form-control-sm" name="heures" placeholder="h/j"></div><div class="col-2"><button type="submit" class="btn btn-success btn-sm w-100">+</button></div></div></form></div></div></div></div><div class="row g-4 mb-4"><div class="col-lg-8"><div class="card shadow-sm"><div class="card-header bg-primary text-white">📈 Évolution</div><div class="card-body"><canvas id="consoChart" height="200"></canvas></div></div></div><div class="col-lg-4"><div class="card shadow-sm"><div class="card-header bg-info text-white">🔬 Régression</div><div class="card-body">{reg_html}</div></div></div></div><div class="card shadow-sm"><div class="card-header bg-dark text-white">📋 Relevés</div><div class="card-body table-responsive"><table class="table table-hover mb-0"><thead><tr><th>Date</th><th>kWh</th><th>Coupure</th><th>Temp.</th><th>Coût</th></tr></thead><tbody>{rel_html}</tbody></table></div></div></div><script>var d={dates_js},c={consos_js};if(d.length>0){{new Chart(document.getElementById('consoChart').getContext('2d'),{{type:'line',data:{{labels:d,datasets:[{{label:'kWh',data:c,borderColor:'#ffd700',tension:0.3,fill:true}}]}}}});}}var al={app_labels},ad={app_data};if(al.length>0){{new Chart(document.getElementById('appChart').getContext('2d'),{{type:'pie',data:{{labels:al,datasets:[{{data:ad}}]}}}});}}</script><script src="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/js/bootstrap.bundle.min.js"></script></body></html>""")

@app.get("/export_excel/{foyer_id}")
async def export_excel(foyer_id: int, db: Session = Depends(get_db)):
    client = db.query(Client).filter(Client.id == foyer_id).first()
    if not client: raise HTTPException(404)
    wb = openpyxl.Workbook()
    ws1 = wb.active; ws1.title = "Relevés"
    hf = Font(bold=True, color="FFFFFF"); hfl = PatternFill(start_color="2c5364", end_color="2c5364", fill_type="solid")
    for i,h in enumerate(["Date","kWh","Coupure(min)","Temp(°C)","Coût(FCFA)"],1):
        c=ws1.cell(row=1,column=i,value=h);c.font=hf;c.fill=hfl
    for i,r in enumerate(list(client.releves),2):
        ws1.cell(row=i,column=1,value=str(r.date_releve))
        ws1.cell(row=i,column=2,value=r.index_compteur)
        ws1.cell(row=i,column=3,value=r.duree_coupure_minutes)
        ws1.cell(row=i,column=4,value=r.temperature_exterieure or 0)
        ws1.cell(row=i,column=5,value=r.cout_estime_fcfa or 0)
    out = io.BytesIO(); wb.save(out); out.seek(0)
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename=WattScope_{client.nom_utilisateur}.xlsx"})

@app.post("/ajouter_appareil")
async def ajouter_appareil(request: Request, foyer_id: int = Form(...), nom_appareil: str = Form(...), puissance_watts: int = Form(...), heures: float = Form(1.0), db: Session = Depends(get_db)):
    db.add(Appareil(foyer_id=foyer_id, nom_appareil=nom_appareil, puissance_watts=puissance_watts, heures_utilisation_jour=heures, nombre_appareils=1))
    db.commit()
    return RedirectResponse(url=f"/dashboard/{foyer_id}?success=✅ Appareil ajouté", status_code=303)
